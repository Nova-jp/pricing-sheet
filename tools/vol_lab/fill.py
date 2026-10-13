"""
VolLab.xlsxのデータシート(3行目以降)だけを書き換える。Excelを使わず、xlsx(zip)の中の
シートのXMLを直接差し替える。グラフ・数式・設定など他の部品には触らない
(openpyxlで読み込んで保存するとグラフが消えるため、この方式にしている)。

書き込むのは数値と日付だけ(文字列は書かない。見出しの1〜2行目は雛形のまま残す)。
開いたときにExcelが全再計算するよう、ブックに fullCalcOnLoad を付ける。
"""

import re
import zipfile
from datetime import date
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Union

from openpyxl.utils import get_column_letter

from tools.template import _sheet_paths

FIRST_ROW = 3  # データシートの見出しは1〜2行目、データは3行目から(tools/template.pyのRECEIVE_AREASと同じ)

Cell = Union[float, int, date, None]
_EXCEL_EPOCH = date(1899, 12, 30)


def _date_style_index(styles_xml: str) -> int:
    """日付の表示形式を持つセルの書式(cellXfsの番号)。雛形の日付セルと同じものを使う。"""
    custom = {int(i): code for i, code in re.findall(r'<numFmt numFmtId="(\d+)" formatCode="([^"]*)"', styles_xml)}
    cell_xfs = re.search(r"<cellXfs[^>]*>(.*?)</cellXfs>", styles_xml, re.S).group(1)
    for n, xf in enumerate(re.findall(r"<xf\b[^>]*?(?:/>|>.*?</xf>)", cell_xfs, re.S)):
        fmt = int(re.search(r'numFmtId="(\d+)"', xf).group(1))
        code = custom.get(fmt, "")
        if fmt in (14, 15, 16, 17, 22) or ("y" in code and "d" in code and "h" not in code):
            return n
    raise ValueError("日付の表示形式の書式がブックにありません(雛形の設定シートの日付セルを確認)")


def _number(value: Union[float, int]) -> str:
    return repr(float(value)) if isinstance(value, float) else str(value)


def _rows_xml(rows: Sequence[Sequence[Cell]], date_style: int) -> str:
    out = []
    for n, row in enumerate(rows):
        r = FIRST_ROW + n
        cells = []
        for c, value in enumerate(row, start=1):
            if value is None:
                continue
            ref = f"{get_column_letter(c)}{r}"
            if isinstance(value, date):
                cells.append(f'<c r="{ref}" s="{date_style}"><v>{(value - _EXCEL_EPOCH).days}</v></c>')
            else:
                cells.append(f'<c r="{ref}"><v>{_number(value)}</v></c>')
        out.append(f'<row r="{r}">{"".join(cells)}</row>')
    return "".join(out)


def replace_sheet_rows(sheet_xml: str, rows: Sequence[Sequence[Cell]], date_style: int) -> str:
    """シートのXMLのうち、FIRST_ROW行目以降を rows に置き換える(それより上の行は残す)。"""
    match = re.search(r"<sheetData\s*/>|<sheetData>(.*?)</sheetData>", sheet_xml, re.S)
    if match is None:
        raise ValueError("シートにsheetDataがありません")
    kept = [m.group(0) for m in re.finditer(r'<row r="(\d+)"[^>]*?(?:/>|>.*?</row>)', match.group(1) or "", re.S)
            if int(m.group(1)) < FIRST_ROW]
    body = "".join(kept) + _rows_xml(rows, date_style)
    xml = sheet_xml[:match.start()] + f"<sheetData>{body}</sheetData>" + sheet_xml[match.end():]
    width = max([len(r) for r in rows] + [1])
    last_row = FIRST_ROW + len(rows) - 1 if rows else FIRST_ROW - 1
    return re.sub(r'<dimension ref="[^"]*"/>', f'<dimension ref="A1:{get_column_letter(width)}{max(last_row, 1)}"/>', xml)


def _full_calc_on_load(workbook_xml: str) -> str:
    if "fullCalcOnLoad" in workbook_xml:
        return re.sub(r'fullCalcOnLoad="[^"]*"', 'fullCalcOnLoad="1"', workbook_xml)
    if "<calcPr" in workbook_xml:
        return re.sub(r"<calcPr\b", '<calcPr fullCalcOnLoad="1"', workbook_xml, count=1)
    return workbook_xml.replace("</workbook>", '<calcPr fullCalcOnLoad="1"/></workbook>')


def fill_workbook(path: Path, tables: Dict[str, List[List[Cell]]]) -> None:
    """tables: {シート名: 行の並び(各行は A列から)}。指定したシートの3行目以降だけを置き換える。"""
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        parts = {info.filename: z.read(info) for info in infos}
    sheet_paths = _sheet_paths(parts)
    missing = set(tables) - set(sheet_paths)
    if missing:
        raise ValueError(f"ブックにデータシートがありません: {sorted(missing)}")
    date_style = _date_style_index(parts["xl/styles.xml"].decode("utf-8"))
    for sheet, rows in tables.items():
        name = sheet_paths[sheet]
        parts[name] = replace_sheet_rows(parts[name].decode("utf-8"), rows, date_style).encode("utf-8")
    parts["xl/workbook.xml"] = _full_calc_on_load(parts["xl/workbook.xml"].decode("utf-8")).encode("utf-8")

    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for info in infos:
            out.writestr(info, parts[info.filename], compress_type=zipfile.ZIP_DEFLATED)
    tmp.replace(path)


def as_rows(dates: Sequence[date], *blocks: Sequence[Sequence[Optional[float]]]) -> List[List[Cell]]:
    """日付の列+値のブロック(同じ行数)を横につないだ行の並び。"""
    return [[d, *[v for block in blocks for v in block[i]]] for i, d in enumerate(dates)]
