"""
ヒストリカル日次カーブDB(data/historical.db)を作るバッチ。朝に手動で実行する。

使い方:
  1. historical_data.xlsx をExcelで開き、LSEGのデータ更新が終わったら保存する
     (このバッチは保存済みの値を読む。Excelで開いたままでもよい)
  2. python -m swap_pricing.curve_batch
     オプション: --xlsx パス / --config パス / --db パス / --rebuild(全日付を作り直す)

処理:
  ① historical_data.xlsx の Mid(日付×テナーのミッド)と Fixing(TONA実績)を読み、
     DBの quotes / fixings に取り込む(同じ日付は置き換え。過去に取り込んだ日付は残る)。
     Volシート(スワップションATMノーマルvol)があれば vol_quotes にも取り込む
     (分析ブックのPCA用。カーブには使わない)
  ② DBにある全日付について、historical_curve.toml で選んだテナーでカーブを構築し、
     暦日ごとのDFを curves に保存する(形式は curve_store.py)。
     - "O/N" はその日のTONA実績(フィキシング対象日=その日)を使う
     - その日に値が無いテナーは外し、あるテナーだけで引く(使用/不使用テナーを記録)
     - 日本の非営業日(ql.Japan)は作らない。LSEGには祝日にも一部テナーの値が
       入ることがある(実データで1テナーだけの日もあった)が、TONAも市場も無い日の
       カーブは意味を持たないため
     - 入力(構築方法・使うレート)のハッシュが前回と同じ日は作り直さない。
       テナー選択や構築方法を変えれば、該当する日が自動で作り直される

カーブの引き方は curve.CURVE_BUILDERS から名前で選ぶだけで、このバッチは引き方に依存しない。
"""

import argparse
import hashlib
import json
import sys
import tomllib
from datetime import date, datetime
from pathlib import Path
from typing import Dict, List, NamedTuple, Optional, Tuple

import openpyxl
import QuantLib as ql

from swap_pricing.conventions import CALENDAR
from swap_pricing.curve import get_curve_builder
from swap_pricing.curve_store import GRID_DAYS, CurveStore, discount_grid
from swap_pricing.paths import HISTORICAL_DB_PATH, REPO_ROOT

DEFAULT_XLSX = REPO_ROOT / "historical_data.xlsx"
DEFAULT_CONFIG = REPO_ROOT / "historical_curve.toml"
ON_TENOR = "O/N"
_DATE_HEADER = "日付"
_VOL_HEADER = "満期x原資産"
VOL_SHEET = "Vol"


class CurveConfig(NamedTuple):
    builder: str
    tenors: List[str]


class BatchSummary(NamedTuple):
    built: List[date]
    unchanged: int
    holidays: List[date]  # 日本の非営業日のため作らなかった日
    with_missing: Dict[date, List[str]]  # テナーが欠けたまま構築した日
    failed: Dict[date, str]


def load_config(path: Path) -> CurveConfig:
    with open(path, "rb") as f:
        raw = tomllib.load(f)
    tenors = [str(t).strip().upper() for t in raw.get("tenors", [])]
    if not tenors:
        raise ValueError(f"{path}: tenors が空です")
    builder = str(raw.get("builder", "")).strip()
    get_curve_builder(builder)  # 未対応の名前はここでエラー
    return CurveConfig(builder=builder, tenors=tenors)


def _as_date(value) -> Optional[date]:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _as_rate(value) -> Optional[float]:
    if value is None or isinstance(value, bool):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None  # 空文字・"#N/A"等のエラー値は欠損扱い


def _rows_after_date_header(ws, header_text: str = _DATE_HEADER):
    """A列が header_text の見出し行を探し、(見出し行, 以降の行のイテレータ) を返す。"""
    rows = ws.iter_rows(values_only=True)
    for row in rows:
        if row and isinstance(row[0], str) and row[0].strip() == header_text:
            return row, rows
    raise ValueError(f"{ws.title}シートに「{header_text}」の見出し行がありません")


def _date_matrix(rows, labels) -> Dict[date, Dict[str, float]]:
    """見出し行の後の行(A列=日付、B列以降=銘柄ごとの値)→ {日付: {銘柄: 値}}。値の無い日は除く。"""
    result: Dict[date, Dict[str, float]] = {}
    for row in rows:
        as_of = _as_date(row[0]) if row else None
        if as_of is None:
            continue  # RIC行・空行
        values = {}
        for label, value in zip(labels, row[1:]):
            v = _as_rate(value)
            if label is not None and v is not None:
                values[label] = v
        if values:
            result[as_of] = values
    return result


def read_vol_sheet(path: Path) -> Dict[date, Dict[str, float]]:
    """historical_data.xlsx の Volシート → {日付: {満期x原資産: ノーマルvol(bp)}}。
    シートが無い(古い雛形)ときは空。銘柄の表記はシートの見出しのまま(例 1Yx10Y)。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if VOL_SHEET not in wb.sheetnames:
            return {}
        header, rows = _rows_after_date_header(wb[VOL_SHEET], _VOL_HEADER)
        labels = [str(h).strip() if h not in (None, "") else None for h in header[1:]]
        return _date_matrix(rows, labels)
    finally:
        wb.close()


def read_historical_workbook(path: Path) -> Tuple[Dict[date, Dict[str, float]], Dict[date, float]]:
    """historical_data.xlsx → ({日付: {テナー: ミッド%}}, {フィキシング対象日: TONA%})。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        header, rows = _rows_after_date_header(wb["Mid"])
        tenors = [str(h).strip().upper() if h not in (None, "") else None for h in header[1:]]
        mid = _date_matrix(rows, tenors)

        _, rows = _rows_after_date_header(wb["Fixing"])
        fixings: Dict[date, float] = {}
        for row in rows:
            d = _as_date(row[0]) if row else None
            rate = _as_rate(row[1]) if row and len(row) > 1 else None
            if d is not None and rate is not None:
                fixings[d] = rate
    finally:
        wb.close()
    if not mid:
        raise ValueError(
            f"{path.name}のMidシートにデータがありません。Excelで開いてLSEGの更新が"
            "終わってから保存し、もう一度実行してください"
        )
    return mid, fixings


def select_quotes(
    quotes: Dict[str, float], fixing: Optional[float], tenors: List[str],
) -> Tuple[Dict[str, float], List[str]]:
    """設定のテナーのうち、その日に値があるものだけを選ぶ → (使うレート, 欠けているテナー)。"""
    selected: Dict[str, float] = {}
    missing: List[str] = []
    for tenor in tenors:
        value = fixing if tenor == ON_TENOR else quotes.get(tenor)
        if value is None:
            missing.append(tenor)
        else:
            selected[tenor] = value
    return selected, missing


def input_hash(builder: str, selected: Dict[str, float]) -> str:
    payload = json.dumps(
        {"builder": builder, "grid_days": GRID_DAYS, "quotes": sorted(selected.items())},
    )
    return hashlib.sha256(payload.encode()).hexdigest()


def run(
    xlsx_path: Path = DEFAULT_XLSX,
    config_path: Path = DEFAULT_CONFIG,
    db_path: Path = HISTORICAL_DB_PATH,
    rebuild: bool = False,
    log=print,
) -> BatchSummary:
    config = load_config(config_path)
    builder = get_curve_builder(config.builder)
    store = CurveStore(db_path)

    mid, fixings = read_historical_workbook(xlsx_path)
    store.save_quotes(mid)
    store.save_fixings(fixings)
    vols = read_vol_sheet(xlsx_path)
    store.save_vol_quotes(vols)
    log(f"取り込み: Mid {len(mid)}日分 / TONA実績 {len(fixings)}日分 / Vol {len(vols)}日分")

    all_quotes = store.quotes()
    all_fixings = store.fixings()
    existing = {} if rebuild else store.input_hashes()

    built: List[date] = []
    unchanged = 0
    holidays: List[date] = []
    with_missing: Dict[date, List[str]] = {}
    failed: Dict[date, str] = {}
    for as_of in sorted(all_quotes):
        if not CALENDAR.isBusinessDay(ql.Date.from_date(as_of)):
            holidays.append(as_of)
            continue
        selected, missing = select_quotes(all_quotes[as_of], all_fixings.get(as_of), config.tenors)
        if not any(t != ON_TENOR for t in selected):
            failed[as_of] = "使えるテナーがありません(O/N以外が全て欠損)"
            continue
        h = input_hash(config.builder, selected)
        if existing.get(as_of) == h:
            unchanged += 1
            continue
        try:
            result = builder(as_of, selected)
            dfs = discount_grid(result.curve, as_of)
        except Exception as e:  # 1日の失敗で全体を止めない(最後にまとめて報告)
            failed[as_of] = str(e)
            continue
        store.save_curve(
            as_of, config.builder, h, result.used_tenors, missing + result.skipped_tenors, dfs,
        )
        built.append(as_of)
        if missing:
            with_missing[as_of] = missing
        if len(built) % 50 == 0:
            log(f"  構築中... {len(built)}日分")

    summary = BatchSummary(built, unchanged, holidays, with_missing, failed)
    log(
        f"カーブ構築: 新規/更新 {len(built)}日 / 変更なし {unchanged}日 / "
        f"非営業日のため除外 {len(holidays)}日 / 失敗 {len(failed)}日"
    )
    if holidays:
        log(f"  非営業日: {', '.join(str(d) for d in holidays)}")
    for d, tenors in sorted(with_missing.items()):
        log(f"  {d}: 欠損テナーを除いて構築 ({', '.join(tenors)})")
    for d, msg in sorted(failed.items()):
        log(f"  {d}: 失敗 {msg}")
    return summary


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="historical_data.xlsx からヒストリカル日次カーブDBを作る")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG)
    parser.add_argument("--db", type=Path, default=HISTORICAL_DB_PATH)
    parser.add_argument("--rebuild", action="store_true", help="入力が同じ日も含めて全日付を作り直す")
    args = parser.parse_args(argv)
    summary = run(args.xlsx, args.config, args.db, args.rebuild)
    return 1 if summary.failed else 0


if __name__ == "__main__":
    sys.exit(main())
