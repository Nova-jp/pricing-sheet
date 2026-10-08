"""
日次カーブDBを作るバッチ(swap_pricing/curve_batch.py)のリグレッションテスト。

LSEGのデータはPublicリポジトリに置けないため、historical_data.xlsx と同じ構成
(Mid: A列「日付」の見出し行+RIC行+日付×テナー / Fixing: 日付・TONA)の合成ブックを
一時フォルダに作って確認する。

実行方法:
    python -m unittest tests.test_curve_batch -v
"""

import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

import openpyxl  # noqa: E402

from swap_pricing import curve_batch  # noqa: E402
from swap_pricing.curve import bootstrap_curve  # noqa: E402
from swap_pricing.curve_store import CurveStore  # noqa: E402
from swap_pricing.historical_pricer import FIXED, historical_swap_rate_series  # noqa: E402
from swap_pricing.swap_pricer import price_swap  # noqa: E402
from tests.test_curve_repricing import SPOT_RATES  # noqa: E402
from tests.test_past_start import CONV, _synthetic_fixings  # noqa: E402

DAYS = [date(2026, 9, 24), date(2026, 9, 25), date(2026, 9, 28)]
TENORS = list(SPOT_RATES)  # Midシートの見出し(12M/24M/36M表記を含む)
FIXINGS = _synthetic_fixings(date(2026, 6, 1), date(2026, 9, 25))  # 9/28分は未公表


def _write_workbook(path: Path, missing=None, mid_days=DAYS):
    """missing: {日付: [欠損させるテナー]}(LSEGのエラー値 "#N/A" を入れる)。"""
    missing = missing or {}
    wb = openpyxl.Workbook()
    mid = wb.active
    mid.title = "Mid"
    mid.append(["ミッドレート ヒストリカル(%)"])
    mid.append(["期間: ..."])
    mid.append([])
    mid.append(["日付"] + TENORS)
    mid.append(["RIC"] + [f"JP{t}ONI=TRDT" for t in TENORS])
    for d in sorted(mid_days, reverse=True):  # LSEGと同じく新しい順
        mid.append([datetime(d.year, d.month, d.day)] + [
            "#N/A" if t in missing.get(d, []) else SPOT_RATES[t] for t in TENORS
        ])
    fixing = wb.create_sheet("Fixing")
    fixing.append(["TONA フィキシング(%)"])
    fixing.append([])
    fixing.append([])
    fixing.append(["日付", "TONA"])
    for d, r in sorted(FIXINGS.items(), reverse=True):
        fixing.append([datetime(d.year, d.month, d.day), r])
    wb.save(path)


def _add_vol_sheet(path: Path):
    """Volシート(historical_data.xlsxと同じ構成: 「満期x原資産」の見出し行+RIC行+新しい順の日付)。"""
    wb = openpyxl.load_workbook(path)
    vol = wb.create_sheet("Vol")
    vol.append(["Implied Vol ヒストリカル(ATMノーマルvol, bp/年)"])
    vol.append(["期間: ..."])
    vol.append([])
    vol.append(["満期x原資産", "1Mx1Y", "1Yx10Y"])
    vol.append(["RIC", "JPYTO1MX1YATM=R", "JPYTO1YX10YATM=R"])
    vol.append([datetime(2026, 9, 25), 44.09, "#N/A"])
    vol.append([datetime(2026, 9, 24), 42.42, 72.5])
    wb.save(path)


def _write_config(path: Path, tenors):
    quoted = ", ".join(f'"{t}"' for t in tenors)
    path.write_text(f'builder = "convex_monotone"\ntenors = [{quoted}]\n', encoding="utf-8")


def _par5y(curve, index, as_of, spot):
    return price_swap(
        curve, as_of, spot, notional=1.0, pay_rec="PAY", tenor="5Y", index=index, **CONV,
    ).target_fixrate


class CurveBatchTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        tmp = Path(self._tmp.name)
        self.xlsx, self.config, self.db = tmp / "historical_data.xlsx", tmp / "c.toml", tmp / "h.db"
        _write_workbook(self.xlsx)
        _write_config(self.config, ["O/N"] + TENORS)

    def tearDown(self):
        self._tmp.cleanup()

    def _run(self, **kwargs):
        return curve_batch.run(self.xlsx, self.config, self.db, log=lambda *_: None, **kwargs)

    def test_builds_curve_equal_to_direct_bootstrap_with_tona_as_on(self):
        summary = self._run()
        self.assertEqual(summary.built, DAYS)
        store = CurveStore(self.db)
        as_of = date(2026, 9, 25)
        quotes = {"O/N": FIXINGS[as_of], **SPOT_RATES}  # O/Nはその日のTONA実績
        boot = bootstrap_curve(as_of, quotes)
        self.assertIn("o/n", boot.used_tenors)
        stored = store.load(as_of)
        self.assertAlmostEqual(
            _par5y(stored.curve, stored.index, as_of, boot.spot_date),
            _par5y(boot.curve, boot.index, as_of, boot.spot_date),
            delta=1e-9,
        )

    def test_missing_on_is_skipped_and_recorded(self):
        """TONA実績が未公表の日(9/28)はO/N抜きで構築し、欠損として記録する。"""
        summary = self._run()
        self.assertEqual(summary.with_missing, {date(2026, 9, 28): ["O/N"]})

    def test_missing_tenor_builds_from_available(self):
        _write_workbook(self.xlsx, missing={date(2026, 9, 25): ["7Y", "15Y"]})
        summary = self._run()
        self.assertIn(date(2026, 9, 25), summary.built)
        self.assertEqual(summary.with_missing[date(2026, 9, 25)], ["7Y", "15Y"])
        as_of = date(2026, 9, 25)
        quotes = {"O/N": FIXINGS[as_of], **{t: r for t, r in SPOT_RATES.items() if t not in ("7Y", "15Y")}}
        boot = bootstrap_curve(as_of, quotes)
        stored = CurveStore(self.db).load(as_of)
        self.assertAlmostEqual(
            _par5y(stored.curve, stored.index, as_of, boot.spot_date),
            _par5y(boot.curve, boot.index, as_of, boot.spot_date),
            delta=1e-9,
        )

    def test_tenor_selection_and_incremental_rebuild(self):
        self._run()
        again = self._run()
        self.assertEqual((again.built, again.unchanged), ([], len(DAYS)))

        _write_config(self.config, [t for t in ["O/N"] + TENORS if t != "40Y"])
        changed = self._run()
        self.assertEqual(changed.built, DAYS)  # テナー選択を変えると作り直される
        self.assertEqual(self._run(rebuild=True).built, DAYS)

    def test_dates_dropped_from_workbook_are_kept_and_rebuilt_from_db(self):
        """xlsxの取得期間から外れた日も、DBの生レートから作り直せる。"""
        self._run()
        _write_workbook(self.xlsx, mid_days=DAYS[1:])  # 9/24がxlsxから外れる
        _write_config(self.config, TENORS)  # O/Nを外す
        summary = self._run()
        # 9/28は元々O/N抜き(TONA未公表)で作っていたので入力が変わらず、作り直さない
        self.assertEqual(summary.built, DAYS[:2])
        self.assertEqual(summary.unchanged, 1)
        self.assertEqual(CurveStore(self.db).dates(), DAYS)

    def test_historical_series_from_db(self):
        """DBのカーブとTONA実績で、過去起算のFIXED系列が計算できる。"""
        self._run()
        store = CurveStore(self.db)
        series = historical_swap_rate_series(
            FIXED, start=date(2026, 7, 1), tenor="2Y", fixings=store.fixings(), curves=store, **CONV,
        )
        self.assertEqual([p.as_of_date for p in series], DAYS)

    def test_japanese_holidays_are_not_built(self):
        """祝日(2026-09-22 国民の休日)にLSEGの値があっても、その日のカーブは作らない。"""
        holiday = date(2026, 9, 22)
        _write_workbook(self.xlsx, mid_days=DAYS + [holiday])
        summary = self._run()
        self.assertEqual(summary.holidays, [holiday])
        self.assertNotIn(holiday, CurveStore(self.db).dates())

    def test_udf_historical_rates_aligns_legs(self):
        """UDF HistoricalRates: 空の行は空欄の列、各脚は1本ずつ計算した値と一致する。"""
        from unittest import mock

        from swap_pricing import udf

        self._run()
        store = CurveStore(self.db)
        conv = ["PA", "act/365fixed", "PA", "act/365fixed", "STD"]
        table = [["", "5y", ""] + conv, ["", "", ""] + ["", "", "", "", ""], ["", "10Y", ""] + conv]
        with mock.patch.object(udf, "_historical_store", lambda: store):
            rows = udf.HistoricalRates(table, "rolling")
            single = udf.HistoricalRate("", "10Y", "", *conv, "ROLLING")
        self.assertEqual(rows[0], ["Date", "A", "B", "C"])
        self.assertEqual([r[0] for r in rows[1:]], DAYS)
        self.assertTrue(all(r[2] == "" for r in rows[1:]))
        self.assertEqual([r[3] for r in rows[1:]], [r[1] for r in single[1:]])

    def test_udf_historical_rates_without_legs_returns_dates(self):
        from unittest import mock

        from swap_pricing import udf

        self._run()
        with mock.patch.object(udf, "_historical_store", lambda: CurveStore(self.db)):
            rows = udf.HistoricalRates([[""] * 8] * 3, "FIXED")
        self.assertEqual(rows[1:], [[d, "", "", ""] for d in DAYS])

    def test_vol_sheet_is_imported(self):
        # Volシートがあれば vol_quotes に取り込む(エラー値は欠損)。無い雛形でもバッチは動く
        self.assertEqual(CurveStore(self.db).vol_quotes() if self.db.exists() else {}, {})
        _add_vol_sheet(self.xlsx)
        self._run()
        self.assertEqual(CurveStore(self.db).vol_quotes(), {
            date(2026, 9, 25): {"1Mx1Y": 44.09},
            date(2026, 9, 24): {"1Mx1Y": 42.42, "1Yx10Y": 72.5},
        })

    def test_empty_workbook_is_rejected(self):
        _write_workbook(self.xlsx, mid_days=[])
        with self.assertRaisesRegex(ValueError, "Midシートにデータがありません"):
            self._run()


if __name__ == "__main__":
    unittest.main()
