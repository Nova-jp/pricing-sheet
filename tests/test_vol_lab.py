"""
tools/vol_lab(ボラ分析のたたき台)のテスト。

    python -m unittest tests.test_vol_lab -v

- data.compute: 使い回しのスワップで読んだFwdが、日付ごとに swap_pricer.par_rate で組み直した値と
  完全に一致すること(満期一定型の変化・累積、オプション追跡型の二乗和・日数、満期までの日数)。
  カーブは日付ごとに水準の違う合成のフラットカーブ(DBは使わない)。
- fill: データシートの3行目以降だけが置き換わり、見出し・他のシートは残ること。
"""

import shutil
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

import QuantLib as ql  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from engine.common.conventions import CALENDAR, forward_start_date  # noqa: E402
from engine.common.curve import spot_date_for  # noqa: E402
from engine.common.historical_store import StoredCurve  # noqa: E402
from engine.common.market_data import Market  # noqa: E402
from engine.common.swap_pricer import par_rate  # noqa: E402
from engine.common.terms import STANDARD, SwapTerms  # noqa: E402
from tools.vol_lab.data import compute, option_expiry  # noqa: E402
from tools.vol_lab.fill import as_rows, fill_workbook  # noqa: E402

EXPIRIES = ["1M", "2M"]
TAILS = ["1Y", "5Y"]
TRACK = ["1M"]


def _business_days(start: date, n: int):
    d = ql.Date.from_date(start)
    out = []
    while len(out) < n:
        if CALENDAR.isBusinessDay(d):
            out.append(d.to_date())
        d += 1
    return out


class _FlatCurves:
    """日付ごとに水準の違うフラットカーブ(HistoricalStoreと同じ dates()/load() を持つ)。"""

    def __init__(self, dates):
        self._dates = dates
        self.rate = {d: 0.01 + 0.0005 * ((n * 7) % 5) for n, d in enumerate(dates)}

    def dates(self):
        return list(self._dates)

    def load(self, as_of_date, until=None):
        curve = ql.FlatForward(ql.Date.from_date(as_of_date), self.rate[as_of_date], ql.Actual365Fixed())
        handle = ql.YieldTermStructureHandle(curve)
        return StoredCurve(curve=handle, index=ql.Tonar(handle))


def _par(curves, as_of, start, tail):
    boot = curves.load(as_of)
    market = Market(as_of, boot.curve, boot.index, {})
    return par_rate(market, SwapTerms(start=start, convention=STANDARD, tenor=tail.lower())).target_fixrate


class ComputeTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dates = _business_days(date(2026, 3, 2), 30)
        cls.curves = _FlatCurves(cls.dates)
        cls.vols = {cls.dates[0]: {"1Mx1Y": 40.0}, cls.dates[-1]: {"2Mx5Y": 55.5}}
        cls.data = compute(cls.curves, cls.vols, expiries=EXPIRIES, tails=TAILS, track_expiries=TRACK)

    def _start(self, i, expiry):
        return forward_start_date(spot_date_for(self.dates[i]), expiry)

    def test_forward_and_daily_change_match_par_rate(self):
        for i in (0, 1, 17, len(self.dates) - 1):
            for e_idx, e in enumerate(EXPIRIES):
                for t_idx, t in enumerate(TAILS):
                    j = e_idx * len(TAILS) + t_idx
                    start = self._start(i, e)
                    self.assertEqual(self.data.fwd[i][j], _par(self.curves, self.dates[i], start, t))
                    if i == 0:
                        self.assertIsNone(self.data.dfwd[i][j])
                    else:
                        # 前日のカーブで「当日の」起算日・満期日を評価した差(ロールダウンを混ぜない)
                        expected = (_par(self.curves, self.dates[i], start, t)
                                    - _par(self.curves, self.dates[i - 1], start, t)) * 100.0
                        self.assertAlmostEqual(self.data.dfwd[i][j], expected, places=12)

    def test_cumulative_sum_of_squares(self):
        for j in range(len(self.data.labels)):
            total = 0.0
            for i, row in enumerate(self.data.dfwd):
                total += row[j] ** 2 if row[j] is not None else 0.0
                self.assertAlmostEqual(self.data.cum_sq[i][j], total, places=9)

    def test_tracking_follows_fixed_dates_until_option_expiry(self):
        r = 0
        start = self._start(r, "1M")
        expiry = option_expiry(start)
        last = max(k for k, d in enumerate(self.dates) if d <= expiry)
        for t_idx, t in enumerate(TAILS):
            values = [_par(self.curves, self.dates[k], start, t) for k in range(r, last + 1)]
            sumsq = sum(((b - a) * 100.0) ** 2 for a, b in zip(values, values[1:]))
            self.assertAlmostEqual(self.data.track_sumsq[r][t_idx], sumsq, places=9)
            self.assertEqual(self.data.track_n[r][t_idx], last - r)
        # 満期日がデータの最終日より後の取引日は追わない
        self.assertIsNone(self.data.track_sumsq[-1][0])
        self.assertIsNone(self.data.track_n[-1][0])

    def test_days_to_expiry_and_iv(self):
        i = 3
        expiry = option_expiry(self._start(i, "2M"))
        self.assertEqual(self.data.days_cal[i][1], (expiry - self.dates[i]).days)
        self.assertEqual(self.data.days_bd[i][1], CALENDAR.businessDaysBetween(
            ql.Date.from_date(self.dates[i]), ql.Date.from_date(expiry), False, True))
        self.assertEqual(self.data.iv[0][0], 40.0)
        self.assertEqual(self.data.iv[-1][3], 55.5)
        self.assertIsNone(self.data.iv[1][0])


class FillTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.path = Path(self.dir) / "book.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "設定"
        ws["A1"] = "残す"
        ws["B2"] = date(2026, 1, 1)
        ws["B2"].number_format = "yyyy-mm-dd"
        data = wb.create_sheet("IV")
        data["A1"], data["A2"], data["B2"] = "説明", "日付", "1Mx1Y"
        data["A3"], data["B3"], data["A9"] = "古い", 1.0, "古い"
        wb.save(self.path)

    def tearDown(self):
        shutil.rmtree(self.dir)

    def test_only_rows_from_third_are_replaced(self):
        rows = as_rows([date(2026, 9, 24), date(2026, 9, 25)], [[None], [44.09]])
        fill_workbook(self.path, {"IV": rows})
        wb = load_workbook(self.path)
        ws = wb["IV"]
        self.assertEqual((ws["A1"].value, ws["A2"].value, ws["B2"].value), ("説明", "日付", "1Mx1Y"))
        self.assertEqual(ws["A3"].value.date(), date(2026, 9, 24))
        self.assertIsNone(ws["B3"].value)
        self.assertEqual(ws["B4"].value, 44.09)
        self.assertIsNone(ws["A9"].value)
        self.assertEqual(wb["設定"]["A1"].value, "残す")
        self.assertTrue(wb.calculation.fullCalcOnLoad)

    def test_unknown_sheet_is_an_error(self):
        with self.assertRaises(ValueError):
            fill_workbook(self.path, {"無い": []})


if __name__ == "__main__":
    unittest.main()
